"""
軌道材料管理システム テスト
Track Material Management System - Tests
"""
import os
import tempfile
from datetime import date

import openpyxl
import pytest

from src.form7 import generate_form7, _fmt_qty, _safe_sheet_name
from src.models import (
    MaterialItem,
    MaterialTransaction,
    ProjectInfo,
    TrackMaterialRegister,
)
from src.register import generate_register


# ─── ヘルパー ───────────────────────────────────────────────


def make_project() -> ProjectInfo:
    return ProjectInfo(
        project_name="テスト工事",
        project_number="TEST-001",
        line_name="テスト線",
        section="A駅〜B駅",
        contractor="テスト工業",
        year=2024,
        month=4,
    )


def make_material(name: str = "50Nレール",
                  previous_balance: float = 10.0) -> MaterialItem:
    item = MaterialItem(
        material_name=name,
        specification="50N 25m",
        unit="本",
        previous_balance=previous_balance,
    )
    item.transactions = [
        MaterialTransaction(
            transaction_date=date(2024, 4, 5),
            description="資材センターより受入",
            received_qty=20,
            issued_qty=0,
        ),
        MaterialTransaction(
            transaction_date=date(2024, 4, 10),
            description="軌道補修",
            received_qty=0,
            issued_qty=8,
        ),
    ]
    return item


def make_register() -> TrackMaterialRegister:
    register = TrackMaterialRegister(project=make_project())
    register.add_material(make_material())
    return register


# ─── モデルテスト ─────────────────────────────────────────────


class TestMaterialItem:
    def test_total_received(self):
        item = make_material(previous_balance=0)
        assert item.total_received == 20

    def test_total_issued(self):
        item = make_material(previous_balance=0)
        assert item.total_issued == 8

    def test_current_balance(self):
        item = make_material(previous_balance=10)
        # 10 (前月) + 20 (受入) - 8 (払出) = 22
        assert item.current_balance == 22

    def test_running_balances(self):
        item = make_material(previous_balance=10)
        balances = item.get_running_balances()
        # After 受入20: 10+20=30; After 払出8: 30-8=22
        assert balances == [30.0, 22.0]

    def test_no_transactions(self):
        item = MaterialItem(
            material_name="テスト材料",
            specification="規格A",
            unit="本",
            previous_balance=5,
        )
        assert item.total_received == 0
        assert item.total_issued == 0
        assert item.current_balance == 5
        assert item.get_running_balances() == []


class TestTrackMaterialRegister:
    def test_add_material(self):
        register = TrackMaterialRegister(project=make_project())
        assert len(register.materials) == 0
        register.add_material(make_material())
        assert len(register.materials) == 1

    def test_validate_consistency_ok(self):
        register = make_register()
        errors = register.validate_consistency()
        assert errors == []

    def test_validate_consistency_negative_balance(self):
        register = TrackMaterialRegister(project=make_project())
        item = MaterialItem(
            material_name="テスト材料",
            specification="規格A",
            unit="本",
            previous_balance=0,
        )
        item.transactions = [
            MaterialTransaction(
                transaction_date=date(2024, 4, 5),
                description="払出超過",
                received_qty=0,
                issued_qty=5,  # 在庫0なのに5本払出
            )
        ]
        register.add_material(item)
        errors = register.validate_consistency()
        assert len(errors) == 1
        assert "テスト材料" in errors[0]


# ─── 帳票7テスト ─────────────────────────────────────────────


class TestForm7:
    def test_generate_creates_file(self):
        register = make_register()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_form7(register, path)
            assert os.path.exists(path)
        finally:
            os.unlink(path)

    def test_one_sheet_per_material(self):
        register = TrackMaterialRegister(project=make_project())
        register.add_material(make_material("材料A"))
        register.add_material(make_material("材料B"))
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_form7(register, path)
            wb = openpyxl.load_workbook(path)
            assert len(wb.sheetnames) == 2
            assert "材料A" in wb.sheetnames
            assert "材料B" in wb.sheetnames
        finally:
            os.unlink(path)

    def test_sheet_contains_material_name(self):
        register = make_register()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_form7(register, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            # 材料情報行（行5）に材料品名が含まれる
            found = False
            for row in ws.iter_rows(values_only=True):
                if row[0] and "50Nレール" in str(row[0]):
                    found = True
                    break
            assert found
        finally:
            os.unlink(path)

    def test_total_row_matches_model(self):
        """合計行の数値がモデルの計算値と一致する"""
        register = make_register()
        material = register.materials[0]
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_form7(register, path)
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            # 合計行を探す
            total_row = None
            for row in ws.iter_rows(values_only=True):
                if row[1] and "合" in str(row[1]):
                    total_row = row
                    break
            assert total_row is not None
            assert str(int(material.total_received)) == str(total_row[2])
            assert str(int(material.total_issued)) == str(total_row[3])
            assert str(int(material.current_balance)) == str(total_row[4])
        finally:
            os.unlink(path)


# ─── 軌道材料管理簿テスト ────────────────────────────────────


class TestRegister:
    def test_generate_creates_file(self):
        register = make_register()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_register(register, path)
            assert os.path.exists(path)
        finally:
            os.unlink(path)

    def test_sheet_name(self):
        register = make_register()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_register(register, path)
            wb = openpyxl.load_workbook(path)
            assert "軌道材料管理簿" in wb.sheetnames
        finally:
            os.unlink(path)

    def test_material_data_in_sheet(self):
        register = make_register()
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            path = f.name
        try:
            generate_register(register, path)
            wb = openpyxl.load_workbook(path)
            ws = wb["軌道材料管理簿"]
            found = False
            for row in ws.iter_rows(values_only=True):
                if row[1] and "50Nレール" in str(row[1]):
                    found = True
                    break
            assert found
        finally:
            os.unlink(path)

    def test_totals_consistent_with_form7(self):
        """
        軌道材料管理簿の合計行と帳票7の合計行が整合していることを確認する。
        同じ Register データから生成された両帳票の当月受入・払出が一致すること。
        """
        register = make_register()
        material = register.materials[0]

        with tempfile.TemporaryDirectory() as tmpdir:
            reg_path = os.path.join(tmpdir, "register.xlsx")
            f7_path = os.path.join(tmpdir, "form7.xlsx")

            generate_register(register, reg_path)
            generate_form7(register, f7_path)

            # 管理簿の当月受入・払出を取得
            wb_reg = openpyxl.load_workbook(reg_path)
            ws_reg = wb_reg["軌道材料管理簿"]
            reg_received = reg_issued = None
            for row in ws_reg.iter_rows(values_only=True):
                if row[1] and "50Nレール" in str(row[1]):
                    reg_received = float(row[5])
                    reg_issued = float(row[6])
                    break

            # 帳票7の合計受入・払出を取得
            wb_f7 = openpyxl.load_workbook(f7_path)
            ws_f7 = wb_f7["50Nレール"]
            f7_received = f7_issued = None
            for row in ws_f7.iter_rows(values_only=True):
                if row[1] and "合" in str(row[1]):
                    f7_received = float(row[2]) if row[2] else 0
                    f7_issued = float(row[3]) if row[3] else 0
                    break

            # 両帳票の整合性確認
            assert reg_received is not None, "管理簿の受入数量が見つかりません"
            assert f7_received is not None, "帳票7の合計行が見つかりません"
            assert reg_received == f7_received, (
                f"受入数量不一致: 管理簿={reg_received}, 帳票7={f7_received}"
            )
            assert reg_issued == f7_issued, (
                f"払出数量不一致: 管理簿={reg_issued}, 帳票7={f7_issued}"
            )


# ─── ユーティリティテスト ─────────────────────────────────────


class TestUtils:
    def test_fmt_qty_integer(self):
        assert _fmt_qty(5.0) == "5"
        assert _fmt_qty(100.0) == "100"

    def test_fmt_qty_decimal(self):
        assert _fmt_qty(5.5) == "5.5"
        assert _fmt_qty(1.250) == "1.25"

    def test_safe_sheet_name_normal(self):
        assert _safe_sheet_name("材料A") == "材料A"

    def test_safe_sheet_name_invalid_chars(self):
        result = _safe_sheet_name("材料:A/B")
        assert ":" not in result
        assert "/" not in result

    def test_safe_sheet_name_max_length(self):
        long_name = "あ" * 50
        result = _safe_sheet_name(long_name)
        assert len(result) <= 31
