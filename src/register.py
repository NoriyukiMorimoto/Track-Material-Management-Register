"""
軌道材料管理簿 生成モジュール
Track Material Management Register Generator
"""
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from .models import TrackMaterialRegister

# 色定義
HEADER_FILL_COLOR = "1F4E79"   # 濃紺（タイトル背景）
SUBHEADER_FILL_COLOR = "2E75B6"
COLUMN_FILL_COLOR = "D6E4F0"   # 薄水色（列ヘッダー背景）
TOTAL_FILL_COLOR = "FFF2CC"    # 薄黄（合計行背景）
INFO_FILL_COLOR = "EBF5FB"     # 情報行背景
WHITE = "FFFFFF"
BLACK = "000000"

# 列定義
COL_NO = 1              # No.
COL_NAME = 2            # 材料品名
COL_SPEC = 3            # 規格・寸法
COL_UNIT = 4            # 単位
COL_PREV_BAL = 5        # 前月繰越
COL_RECEIVED = 6        # 当月受入
COL_ISSUED = 7          # 当月払出
COL_CURR_BAL = 8        # 翌月繰越
COL_REMARKS = 9         # 備考

TOTAL_COLS = 9

# 列幅
COL_WIDTHS = {
    COL_NO: 6,
    COL_NAME: 22,
    COL_SPEC: 20,
    COL_UNIT: 8,
    COL_PREV_BAL: 14,
    COL_RECEIVED: 14,
    COL_ISSUED: 14,
    COL_CURR_BAL: 14,
    COL_REMARKS: 20,
}

# 行高さ
ROW_HEIGHT_TITLE = 32
ROW_HEIGHT_INFO = 22
ROW_HEIGHT_COL_HEADER = 24
ROW_HEIGHT_DATA = 20
ROW_HEIGHT_TOTAL = 22

# フォント定義
FONT_TITLE = Font(name="MS ゴシック", size=16, bold=True, color=WHITE)
FONT_INFO = Font(name="MS ゴシック", size=10, color=BLACK)
FONT_COL_HEADER = Font(name="MS ゴシック", size=10, bold=True, color=BLACK)
FONT_DATA = Font(name="MS 明朝", size=10, color=BLACK)
FONT_TOTAL = Font(name="MS ゴシック", size=10, bold=True, color=BLACK)

# 罫線定義
THIN = Side(border_style="thin", color=BLACK)
MEDIUM = Side(border_style="medium", color=BLACK)
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _make_fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _set_col_widths(ws) -> None:
    for col_idx, width in COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _apply_outer_border(ws, min_row: int, max_row: int,
                        min_col: int, max_col: int) -> None:
    """指定範囲の外枠に太罫線を引く"""
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                             min_col=min_col, max_col=max_col):
        for cell in row:
            left = MEDIUM if cell.column == min_col else THIN
            right = MEDIUM if cell.column == max_col else THIN
            top = MEDIUM if cell.row == min_row else THIN
            bottom = MEDIUM if cell.row == max_row else THIN
            cell.border = Border(left=left, right=right,
                                  top=top, bottom=bottom)


def _fmt_qty(value: float) -> str:
    """数量を整数/小数で適切に表示する"""
    if value == int(value):
        return str(int(value))
    return f"{value:.3f}".rstrip("0").rstrip(".")


def _data_alignment(col: int) -> Alignment:
    """列に応じた文字寄せを返す"""
    if col in (COL_PREV_BAL, COL_RECEIVED, COL_ISSUED, COL_CURR_BAL):
        return Alignment(horizontal="right", vertical="center", indent=1)
    if col == COL_NO:
        return Alignment(horizontal="center", vertical="center")
    if col == COL_UNIT:
        return Alignment(horizontal="center", vertical="center")
    return Alignment(horizontal="left", vertical="center", indent=1)


def generate_register(register: TrackMaterialRegister,
                      output_path: str) -> None:
    """
    軌道材料管理簿を生成する。

    Args:
        register: 軌道材料管理簿データ
        output_path: 出力先Excelファイルパス
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "軌道材料管理簿"

    _set_col_widths(ws)

    project = register.project

    # ── タイトル行 ──────────────────────────────
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=TOTAL_COLS)
    title_cell = ws.cell(row=1, column=1, value="軌道材料管理簿")
    title_cell.font = FONT_TITLE
    title_cell.fill = _make_fill(HEADER_FILL_COLOR)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = ROW_HEIGHT_TITLE

    # ── プロジェクト情報 ─────────────────────────
    # 行2: 工事名 | 工事番号
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=9)
    _write_info_cell(ws, 2, 1, f"工事名：{project.project_name}")
    _write_info_cell(ws, 2, 6, f"工事番号：{project.project_number}")

    # 行3: 路線名 | 工区・区間
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=5)
    ws.merge_cells(start_row=3, start_column=6, end_row=3, end_column=9)
    _write_info_cell(ws, 3, 1, f"路線名：{project.line_name}")
    _write_info_cell(ws, 3, 6, f"工区・区間：{project.section}")

    # 行4: 施工業者 | 年月
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=5)
    ws.merge_cells(start_row=4, start_column=6, end_row=4, end_column=9)
    _write_info_cell(ws, 4, 1, f"施工業者：{project.contractor}")
    _write_info_cell(ws, 4, 6,
                     f"{project.year}年{project.month:02d}月分")

    for row in (2, 3, 4):
        ws.row_dimensions[row].height = ROW_HEIGHT_INFO

    # ── 列ヘッダー ───────────────────────────────
    header_row = 5
    col_headers = [
        "No.", "材料品名", "規格・寸法", "単位",
        "前月繰越", "当月受入", "当月払出", "翌月繰越", "備考",
    ]
    for col_idx, header in enumerate(col_headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = FONT_COL_HEADER
        cell.fill = _make_fill(COLUMN_FILL_COLOR)
        cell.alignment = Alignment(horizontal="center",
                                   vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = ROW_HEIGHT_COL_HEADER

    # ── データ行 ────────────────────────────────
    data_start_row = header_row + 1
    total_received = 0.0
    total_issued = 0.0

    for i, material in enumerate(register.materials):
        row = data_start_row + i
        ws.cell(row=row, column=COL_NO, value=i + 1)
        ws.cell(row=row, column=COL_NAME, value=material.material_name)
        ws.cell(row=row, column=COL_SPEC, value=material.specification)
        ws.cell(row=row, column=COL_UNIT, value=material.unit)
        ws.cell(row=row, column=COL_PREV_BAL,
                value=_fmt_qty(material.previous_balance))
        ws.cell(row=row, column=COL_RECEIVED,
                value=_fmt_qty(material.total_received))
        ws.cell(row=row, column=COL_ISSUED,
                value=_fmt_qty(material.total_issued))
        ws.cell(row=row, column=COL_CURR_BAL,
                value=_fmt_qty(material.current_balance))
        ws.cell(row=row, column=COL_REMARKS, value="")

        for col in range(1, TOTAL_COLS + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = FONT_DATA
            cell.fill = _make_fill(WHITE)
            cell.alignment = _data_alignment(col)
            cell.border = THIN_BORDER
        ws.row_dimensions[row].height = ROW_HEIGHT_DATA

        total_received += material.total_received
        total_issued += material.total_issued

    # ── 合計行 ──────────────────────────────────
    total_row = data_start_row + len(register.materials)
    ws.merge_cells(start_row=total_row, start_column=1,
                   end_row=total_row, end_column=4)
    ws.cell(row=total_row, column=1, value="合　計")
    ws.cell(row=total_row, column=COL_PREV_BAL, value="")
    ws.cell(row=total_row, column=COL_RECEIVED,
            value=_fmt_qty(total_received))
    ws.cell(row=total_row, column=COL_ISSUED,
            value=_fmt_qty(total_issued))
    ws.cell(row=total_row, column=COL_CURR_BAL, value="")
    ws.cell(row=total_row, column=COL_REMARKS, value="")

    for col in range(1, TOTAL_COLS + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.font = FONT_TOTAL
        cell.fill = _make_fill(TOTAL_FILL_COLOR)
        cell.alignment = Alignment(horizontal="center"
                                   if col <= 4 else _data_alignment(col).horizontal,
                                   vertical="center", indent=1)
        cell.border = THIN_BORDER
    ws.row_dimensions[total_row].height = ROW_HEIGHT_TOTAL

    # 外枠太罫線
    _apply_outer_border(ws, header_row, total_row, 1, TOTAL_COLS)

    # ── 確認・承認欄 ────────────────────────────
    sign_row = total_row + 2
    _write_signature_section(ws, sign_row)

    wb.save(output_path)


def _write_info_cell(ws, row: int, col: int, value: str) -> None:
    """プロジェクト情報セルを書き込む"""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = FONT_INFO
    cell.fill = _make_fill(INFO_FILL_COLOR)
    cell.alignment = Alignment(horizontal="left",
                               vertical="center", indent=1)


def _write_signature_section(ws, start_row: int) -> None:
    """確認・承認欄を書き込む"""
    labels = ["作成", "確認", "承認"]
    col_groups = [(1, 3), (4, 6), (7, 9)]

    ws.merge_cells(start_row=start_row, start_column=1,
                   end_row=start_row, end_column=9)
    header = ws.cell(row=start_row, column=1, value="確認・承認")
    header.font = Font(name="MS ゴシック", size=10, bold=True, color=WHITE)
    header.fill = _make_fill(SUBHEADER_FILL_COLOR)
    header.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[start_row].height = ROW_HEIGHT_INFO

    for label, (start_col, end_col) in zip(labels, col_groups):
        label_row = start_row + 1
        sign_row = start_row + 2

        ws.merge_cells(start_row=label_row, start_column=start_col,
                       end_row=label_row, end_column=end_col)
        lbl_cell = ws.cell(row=label_row, column=start_col, value=label)
        lbl_cell.font = FONT_COL_HEADER
        lbl_cell.fill = _make_fill(COLUMN_FILL_COLOR)
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center")
        lbl_cell.border = THIN_BORDER
        ws.row_dimensions[label_row].height = ROW_HEIGHT_INFO

        ws.merge_cells(start_row=sign_row, start_column=start_col,
                       end_row=sign_row, end_column=end_col)
        sign_cell = ws.cell(row=sign_row, column=start_col, value="")
        sign_cell.font = FONT_DATA
        sign_cell.fill = _make_fill(WHITE)
        sign_cell.border = THIN_BORDER
        ws.row_dimensions[sign_row].height = 36
