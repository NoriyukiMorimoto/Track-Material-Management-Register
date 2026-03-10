"""
帳票7 軌道材料受払簿 生成モジュール
Form 7 - Track Material Receipt/Payment Record Generator
"""
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

from .models import MaterialItem, TrackMaterialRegister

# 色定義
HEADER_FILL_COLOR = "1F4E79"     # 濃紺（ヘッダー背景）
SUBHEADER_FILL_COLOR = "2E75B6"  # 中紺（サブヘッダー背景）
COLUMN_FILL_COLOR = "D6E4F0"     # 薄水色（列ヘッダー背景）
TOTAL_FILL_COLOR = "FCE4D6"      # 薄橙（合計行背景）
WHITE = "FFFFFF"
BLACK = "000000"

# 列定義
COL_DATE = 1        # 月日
COL_DESC = 2        # 摘要
COL_RECEIVED = 3    # 受入数量
COL_ISSUED = 4      # 払出数量
COL_BALANCE = 5     # 残高
COL_REMARKS = 6     # 備考

# 列幅
COL_WIDTHS = {
    COL_DATE: 12,
    COL_DESC: 28,
    COL_RECEIVED: 14,
    COL_ISSUED: 14,
    COL_BALANCE: 14,
    COL_REMARKS: 22,
}

# 行高さ
ROW_HEIGHT_HEADER = 28
ROW_HEIGHT_SUBHEADER = 22
ROW_HEIGHT_DATA = 18
ROW_HEIGHT_TOTAL = 20

# フォント定義
FONT_TITLE = Font(name="MS ゴシック", size=14, bold=True, color=WHITE)
FONT_SUBHEADER = Font(name="MS ゴシック", size=10, bold=True, color=WHITE)
FONT_COL_HEADER = Font(name="MS ゴシック", size=10, bold=True, color=BLACK)
FONT_DATA = Font(name="MS 明朝", size=10, color=BLACK)
FONT_TOTAL = Font(name="MS ゴシック", size=10, bold=True, color=BLACK)

# 罫線定義
THIN = Side(border_style="thin", color=BLACK)
MEDIUM = Side(border_style="medium", color=BLACK)
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MEDIUM_BORDER = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
OUTER_MEDIUM = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=THIN)
OUTER_BOTTOM = Border(left=MEDIUM, right=MEDIUM, top=THIN, bottom=MEDIUM)


def _make_fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _set_col_widths(ws) -> None:
    for col_idx, width in COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _apply_outer_border(ws, min_row: int, max_row: int,
                        min_col: int, max_col: int) -> None:
    """指定範囲の外枠に太罫線を引く"""
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                             min_col=min_col, max_col=max_col):
        for cell in row:
            left = MEDIUM if cell.column == min_col else THIN
            right = MEDIUM if cell.column == max_col else THIN
            top = MEDIUM if cell.row == min_row else THIN
            bottom = MEDIUM if cell.row == max_row else THIN
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def _write_form7_sheet(ws, material: MaterialItem, project_info) -> None:
    """
    1枚の帳票7シートに材料受払データを書き込む
    """
    _set_col_widths(ws)

    # ── タイトル行 ──────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    title_cell = ws.cell(row=1, column=1,
                         value="帳票7　軌道材料受払簿")
    title_cell.font = FONT_TITLE
    title_cell.fill = _make_fill(HEADER_FILL_COLOR)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = ROW_HEIGHT_HEADER

    # ── プロジェクト情報 ─────────────────────────
    info_rows = [
        (f"工事名：{project_info.project_name}",
         f"工事番号：{project_info.project_number}"),
        (f"路線名：{project_info.line_name}",
         f"工区・区間：{project_info.section}"),
        (f"施工業者：{project_info.contractor}",
         f"{project_info.year}年{project_info.month:02d}月分"),
    ]
    for i, (left_text, right_text) in enumerate(info_rows):
        row = i + 2
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=3)
        ws.merge_cells(start_row=row, start_column=4,
                       end_row=row, end_column=6)
        cell_l = ws.cell(row=row, column=1, value=left_text)
        cell_r = ws.cell(row=row, column=4, value=right_text)
        for cell in (cell_l, cell_r):
            cell.font = Font(name="MS ゴシック", size=10, color=BLACK)
            cell.fill = _make_fill("EBF5FB")
            cell.alignment = Alignment(horizontal="left",
                                       vertical="center", indent=1)
        ws.row_dimensions[row].height = ROW_HEIGHT_SUBHEADER

    # ── 材料情報 ─────────────────────────────────
    mat_row = 5
    mat_info = (
        f"材料品名：{material.material_name}　　"
        f"規格：{material.specification}　　"
        f"単位：{material.unit}"
    )
    ws.merge_cells(start_row=mat_row, start_column=1,
                   end_row=mat_row, end_column=6)
    mat_cell = ws.cell(row=mat_row, column=1, value=mat_info)
    mat_cell.font = Font(name="MS ゴシック", size=10, bold=True, color=WHITE)
    mat_cell.fill = _make_fill(SUBHEADER_FILL_COLOR)
    mat_cell.alignment = Alignment(horizontal="left",
                                   vertical="center", indent=1)
    ws.row_dimensions[mat_row].height = ROW_HEIGHT_SUBHEADER

    # ── 列ヘッダー ───────────────────────────────
    col_headers = ["月日", "摘要", "受入数量", "払出数量", "残高", "備考"]
    header_row = 6
    for col_idx, header in enumerate(col_headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = FONT_COL_HEADER
        cell.fill = _make_fill(COLUMN_FILL_COLOR)
        cell.alignment = Alignment(horizontal="center",
                                   vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[header_row].height = ROW_HEIGHT_SUBHEADER

    # ── 前月繰越行 ──────────────────────────────
    carry_row = 7
    running_balance = material.previous_balance
    ws.cell(row=carry_row, column=COL_DATE, value="")
    ws.cell(row=carry_row, column=COL_DESC, value="前月繰越")
    ws.cell(row=carry_row, column=COL_RECEIVED, value="")
    ws.cell(row=carry_row, column=COL_ISSUED, value="")
    ws.cell(row=carry_row, column=COL_BALANCE,
            value=_fmt_qty(running_balance))
    ws.cell(row=carry_row, column=COL_REMARKS, value="")
    for col in range(1, 7):
        cell = ws.cell(row=carry_row, column=col)
        cell.font = Font(name="MS 明朝", size=10, italic=True, color="595959")
        cell.fill = _make_fill("F2F2F2")
        cell.alignment = _data_alignment(col)
        cell.border = THIN_BORDER
    ws.row_dimensions[carry_row].height = ROW_HEIGHT_DATA

    # ── 取引データ行 ──────────────────────────────
    running_balances = material.get_running_balances()
    data_start_row = carry_row + 1
    for i, (trans, balance) in enumerate(
            zip(material.transactions, running_balances)):
        row = data_start_row + i
        date_str = f"{trans.transaction_date.month}/{trans.transaction_date.day}"
        ws.cell(row=row, column=COL_DATE, value=date_str)
        ws.cell(row=row, column=COL_DESC, value=trans.description)
        ws.cell(row=row, column=COL_RECEIVED,
                value=_fmt_qty(trans.received_qty) if trans.received_qty else "")
        ws.cell(row=row, column=COL_ISSUED,
                value=_fmt_qty(trans.issued_qty) if trans.issued_qty else "")
        ws.cell(row=row, column=COL_BALANCE, value=_fmt_qty(balance))
        ws.cell(row=row, column=COL_REMARKS, value=trans.remarks)
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.font = FONT_DATA
            cell.fill = _make_fill(WHITE)
            cell.alignment = _data_alignment(col)
            cell.border = THIN_BORDER
        ws.row_dimensions[row].height = ROW_HEIGHT_DATA

    # ── 合計行 ──────────────────────────────────
    total_row = data_start_row + len(material.transactions)
    ws.cell(row=total_row, column=COL_DATE, value="")
    ws.cell(row=total_row, column=COL_DESC, value="合　計")
    ws.cell(row=total_row, column=COL_RECEIVED,
            value=_fmt_qty(material.total_received))
    ws.cell(row=total_row, column=COL_ISSUED,
            value=_fmt_qty(material.total_issued))
    ws.cell(row=total_row, column=COL_BALANCE,
            value=_fmt_qty(material.current_balance))
    ws.cell(row=total_row, column=COL_REMARKS, value="翌月繰越")
    for col in range(1, 7):
        cell = ws.cell(row=total_row, column=col)
        cell.font = FONT_TOTAL
        cell.fill = _make_fill(TOTAL_FILL_COLOR)
        cell.alignment = _data_alignment(col)
        cell.border = THIN_BORDER
    ws.row_dimensions[total_row].height = ROW_HEIGHT_TOTAL

    # 外枠太罫線
    last_data_row = total_row
    _apply_outer_border(ws, header_row, last_data_row, 1, 6)


def _fmt_qty(value: float) -> str:
    """数量を整数/小数で適切に表示する"""
    if value == int(value):
        return str(int(value))
    return f"{value:.3f}".rstrip("0").rstrip(".")


def _data_alignment(col: int) -> Alignment:
    """列に応じた文字寄せを返す"""
    if col in (COL_RECEIVED, COL_ISSUED, COL_BALANCE):
        return Alignment(horizontal="right", vertical="center", indent=1)
    if col == COL_DATE:
        return Alignment(horizontal="center", vertical="center")
    return Alignment(horizontal="left", vertical="center", indent=1)


def generate_form7(register: TrackMaterialRegister, output_path: str) -> None:
    """
    帳票7（軌道材料受払簿）を生成する。
    材料品目ごとに1シートを作成する。

    Args:
        register: 軌道材料管理簿データ
        output_path: 出力先Excelファイルパス
    """
    wb = Workbook()
    wb.remove(wb.active)  # デフォルトシートを削除

    for material in register.materials:
        # シート名は材料品名（最大31文字、Excelの制限）
        sheet_name = _safe_sheet_name(material.material_name)
        ws = wb.create_sheet(title=sheet_name)
        _write_form7_sheet(ws, material, register.project)

    wb.save(output_path)


def _safe_sheet_name(name: str) -> str:
    """Excelシート名として使用できない文字を除去し、31文字以内にする"""
    invalid_chars = r'\/*?:[]'
    for ch in invalid_chars:
        name = name.replace(ch, "")
    return name[:31]
